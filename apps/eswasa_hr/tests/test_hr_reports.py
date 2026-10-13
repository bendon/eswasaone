"""HR reports and export (additions plan 07) — pure rules, plus run and export on a fake frappe."""

from __future__ import annotations

import base64
import importlib.util
import unittest
from datetime import date

from eswasa_hr import hr_reports as hr
from eswasa_hr import sensitive
from tests import _fakefrappe as ff

TODAY = date(2026, 10, 16)


class RuleTests(unittest.TestCase):
    def test_no_report_outputs_a_protected_field(self):
        protected = set(sensitive.protected_labels())
        for r in hr.REPORTS.values():
            hr.no_sensitive(r["columns"])
            self.assertFalse({c["key"] for c in r["columns"]} & protected, r["id"])
        for key in ("date_of_birth", "bank_ac_no", "ctc", "custom_medical_aid_number", "custom_dependants"):
            with self.assertRaises(ValueError):
                hr.no_sensitive([hr.col("employee", "ID"), hr.col(key, "Leak")])

    def test_disciplinary_register_is_counts_only(self):
        r = hr.REPORTS["disciplinary"]
        keys = {c["key"] for c in r["columns"]}
        self.assertFalse(keys & set(hr.NAME_KEYS))
        self.assertFalse(keys & {"allegation", "outcome", "employee_name"})
        self.assertFalse(hr.has_names(r["columns"]))
        self.assertEqual(r["roles"], hr.MANAGER_ROLES)

    def test_roles(self):
        reports = hr.REPORTS
        self.assertTrue(hr.allowed(reports["staff-list"], ["HR User"]))
        self.assertFalse(hr.allowed(reports["disciplinary"], ["HR User"]))
        self.assertTrue(hr.allowed(reports["disciplinary"], ["HR Manager"]))
        self.assertTrue(hr.allowed(reports["staff-list"], ["Employee"], has_reports=True))
        self.assertFalse(hr.allowed(reports["staff-list"], ["Employee"], has_reports=False))
        self.assertFalse(hr.allowed(reports["leave-owed"], ["Employee"], has_reports=True))
        team = sorted(r["id"] for r in reports.values() if r["team_scoped"])
        self.assertEqual(team, ["expiring", "leave-taken", "staff-list"])

    def test_scope_keeps_only_the_managers_reports(self):
        rows = [{"employee": "E1"}, {"employee": "E2"}, {"employee": None}]
        self.assertEqual(hr.apply_scope(rows, hr.Scope()), rows)
        self.assertEqual(hr.apply_scope(rows, None), rows)
        self.assertEqual(hr.apply_scope(rows, hr.Scope(people=["E2"], label="team")), [{"employee": "E2"}])

    def test_csv_quoting_and_formula_guard(self):
        cols = [hr.col("a", "Name"), hr.col("b", "Note"), hr.col("n", "Days", "number")]
        rows = [{"a": 'Dlamini, "Sipho"', "b": "=HYPERLINK(\"x\")", "n": 2.0}, {"a": "+268", "b": "line\nbreak", "n": -1}, {"a": "@x", "b": "-3", "n": None}]
        text = hr.to_csv(cols, rows).decode("utf-8")
        self.assertTrue(text.startswith("﻿Name,Note,Days\r\n"))
        lines = text[1:].split("\r\n")
        self.assertEqual(lines[1], '"Dlamini, ""Sipho""","\'=HYPERLINK(""x"")",2')
        self.assertTrue(lines[2].startswith("'+268,\"line\nbreak\","))
        self.assertTrue(text.rstrip("\r\n").endswith("'@x,'-3,"))
        self.assertEqual(hr.safe_cell(5), 5)
        self.assertEqual(hr.safe_cell("\tx"), "'\tx")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl not installed here (Frappe ships it)")
    def test_xlsx_never_stores_a_formula(self):
        import io

        from openpyxl import load_workbook

        data = hr.to_xlsx([hr.col("a", "A"), hr.col("d", "D", "date")], [{"a": "=1+1", "d": "2026-10-16"}], "Staff list")
        ws = load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws["A2"].value, "'=1+1")
        self.assertEqual(ws.title, "Staff list")

    def test_filters_defaults_and_errors(self):
        r = hr.REPORTS["leave-taken"]
        values, errs = hr.clean_filters(r, None, TODAY)
        self.assertEqual(errs, [])
        self.assertEqual(values, {"from_date": "2026-01-01", "to_date": "2026-10-16"})
        values, errs = hr.clean_filters(r, '{"from_date": "2026-11-01", "to_date": "2026-10-01", "nope": 1}', TODAY)
        self.assertIn("The start date is after the end date.", errs)
        self.assertNotIn("nope", values)
        _v, errs = hr.clean_filters(hr.REPORTS["expiring"], {"days": "soon"}, TODAY)
        self.assertEqual(errs, ["Within (days) must be a whole number."])
        _v, errs = hr.clean_filters(hr.REPORTS["staff-list"], {"status": "Retired"}, TODAY)
        self.assertEqual(errs, ["Choose status from the list."])

    def test_export_needs_a_reason_only_when_it_names_people(self):
        self.assertIn("names people", hr.export_error(hr.REPORTS["staff-list"], "short", "csv"))
        self.assertIsNone(hr.export_error(hr.REPORTS["staff-list"], "Audit of department records", "xlsx"))
        self.assertIsNone(hr.export_error(hr.REPORTS["disciplinary"], None, "csv"))
        self.assertEqual(hr.export_error(hr.REPORTS["disciplinary"], None, "pdf"), "Choose CSV or Excel.")
        self.assertEqual(hr.file_name("staff-list", "csv", TODAY), "eswasa-hr-staff-list-2026-10-16.csv")

    def test_bands_counts_and_log_text(self):
        self.assertEqual(hr.service_band(0.4), "Under 1 year")
        self.assertEqual(hr.service_band(5.9), "3 to 5 years")
        self.assertEqual(hr.service_band(25), "21 years or more")
        self.assertIsNone(hr.service_band(None))
        self.assertEqual(hr.years_between("2016-10-16", TODAY), 10.0)
        rows = [{"m": "2026-01", "s": "Verbal warning"}, {"m": "2026-01", "s": "Verbal warning"}, {"m": "2026-02", "s": None}]
        self.assertEqual(hr.count_by(rows, ("m", "s")), [{"m": "2026-01", "s": "Verbal warning", "cases": 2}, {"m": "2026-02", "s": "", "cases": 1}])
        self.assertEqual(hr.log_fields("staff-list", {"b": 1, "a": "x"}, hr.Scope(["E1"], "team")), 'staff-list (team) {"a": "x", "b": 1}')


def _store():
    return {
        "Employee": [
            {"name": "E1", "employee_name": "Ann Boss", "status": "Active", "department": "ICT", "designation": "Manager", "user_id": "boss@x",
             "date_of_joining": "2015-01-01", "date_of_birth": "1980-01-01", "bank_ac_no": "123"},
            {"name": "E2", "employee_name": "Ben Report", "status": "Active", "department": "ICT", "designation": "Officer", "user_id": "ben@x",
             "reports_to": "E1", "date_of_joining": "2024-03-01", "contract_end_date": "2026-11-01"},
            {"name": "E3", "employee_name": "Cee Other", "status": "Active", "department": "Finance", "designation": "Clerk", "user_id": "cee@x",
             "date_of_joining": "2020-06-01"},
        ],
        "Department": [{"name": "ICT", "is_group": 0, "disabled": 0}, {"name": "Finance", "is_group": 0, "disabled": 0}],
        "Disciplinary Case": [
            {"name": "C1", "employee": "E2", "opened_on": "2026-03-04", "offence": "LATE", "sanction": "Verbal warning", "allegation": "Late 5 times"},
            {"name": "C2", "employee": "E3", "opened_on": "2026-03-20", "offence": "LATE", "sanction": "Verbal warning", "allegation": "Late again"},
        ],
        "Disciplinary Offence": [{"name": "LATE", "level": "Minor"}],
    }


DOCTYPES = ("Disciplinary Case", "Disciplinary Offence", "Department", "HR Access Log")


class FlowTests(unittest.TestCase):
    def tearDown(self):
        ff.uninstall()

    def _as(self, user, roles):
        self.store = _store()
        return ff.install(self.store, user=user, roles=roles, today="2026-10-16", doctypes=DOCTYPES)

    def test_catalogue_by_role(self):
        self._as("hr@x", ("HR User",))
        ids = [i["id"] for i in hr.catalogue_impl()["items"]]
        self.assertIn("staff-list", ids)
        self.assertNotIn("disciplinary", ids)
        self._as("boss@x", ("Employee",))
        out = hr.catalogue_impl(team="1")
        self.assertEqual(sorted(i["id"] for i in out["items"]), ["expiring", "leave-taken", "staff-list"])
        self.assertTrue(all(i["scope"] == "team" for i in out["items"]))

    def test_someone_without_reports_is_refused(self):
        f = self._as("cee@x", ("Employee",))
        with self.assertRaises(f.PermissionError):
            hr.run_impl("staff-list")
        with self.assertRaises(f.PermissionError):
            hr.catalogue_impl()
        f = self._as("boss@x", ("Employee",))
        with self.assertRaises(f.PermissionError):
            hr.run_impl("leave-owed")

    def test_manager_sees_only_their_reports(self):
        self._as("boss@x", ("Employee",))
        out = hr.run_impl("staff-list")
        self.assertEqual([r["employee"] for r in out["rows"]], ["E2"])
        self.assertEqual(out["scope"], "team")
        self.assertEqual(out["rows"][0]["reports_to_name"], "Ann Boss")
        self.assertNotIn("date_of_birth", out["rows"][0])
        self.assertNotIn("bank_ac_no", out["rows"][0])

    def test_hr_sees_everyone_and_expiring_contracts(self):
        self._as("hr@x", ("HR User",))
        self.assertEqual(hr.run_impl("staff-list")["total"], 3)
        rows = hr.run_impl("expiring", {"kind": "Contract"})["rows"]
        self.assertEqual([(r["employee"], r["days_left"]) for r in rows], [("E2", 16)])

    def test_disciplinary_counts_without_names_for_hr_managers_only(self):
        f = self._as("hr@x", ("HR User",))
        with self.assertRaises(f.PermissionError):
            hr.run_impl("disciplinary")
        self._as("hrm@x", ("HR Manager",))
        out = hr.run_impl("disciplinary")
        self.assertEqual(out["rows"], [{"month": "2026-03", "level": "Minor", "sanction": "Verbal warning", "cases": 2}])
        self.assertNotIn("Late", str(out))

    def test_export_logs_first_then_builds_the_file(self):
        self._as("hrm@x", ("HR Manager",))
        seen = []
        original = hr.REPORTS["staff-list"]["rows"]

        def spy(f, scope):
            seen.append(len(self.store.get("HR Access Log", [])))
            return original(f, scope)

        hr.REPORTS["staff-list"]["rows"] = spy
        try:
            out = hr.export_impl("staff-list", {"department": "ICT"}, "csv", "Annual staff audit for the board")
        finally:
            hr.REPORTS["staff-list"]["rows"] = original
        self.assertEqual(seen, [1])
        log = self.store["HR Access Log"][0]
        self.assertEqual(log["level"], "Report export")
        self.assertEqual(log["viewer"], "hrm@x")
        self.assertEqual(log["fields"], 'staff-list {"department": "ICT", "status": "Active"}')
        self.assertEqual(log["reason"], "Annual staff audit for the board")
        text = base64.b64decode(out["content"]).decode("utf-8")
        self.assertIn("Ann Boss", text)
        self.assertNotIn("Cee Other", text)
        self.assertEqual(out["rows"], 2)
        self.assertEqual(out["content_type"], "text/csv; charset=utf-8")

    def test_export_refused_without_reason_writes_no_log(self):
        f = self._as("hrm@x", ("HR Manager",))
        with self.assertRaises(f.ValidationError):
            hr.export_impl("staff-list", None, "csv", "")
        self.assertEqual(self.store.get("HR Access Log", []), [])
        hr.export_impl("disciplinary", None, "csv", None)
        self.assertEqual(self.store["HR Access Log"][0]["reason"], "Counts only, no names")

    def test_whitelisted_export_needs_confirm(self):
        f = self._as("hrm@x", ("HR Manager",))
        with self.assertRaises(f.ValidationError):
            hr.export("disciplinary", None, "csv", None, confirm=False)


if __name__ == "__main__":
    unittest.main()
